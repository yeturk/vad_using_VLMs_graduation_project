"""Build the project presentation (light / projector-friendly theme, English).

Run from the project root with a Python that has python-pptx + Pillow:
    python presentation/build_pptx.py
Output: presentation/VAD_Presentation.pptx
"""
import os
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "trailer" / "trailer_assets"
OUT = Path(os.environ.get("DECK_OUT", ROOT / "presentation" / "VAD_Presentation.pptx"))

# ---- palette (light theme) -------------------------------------------------
INK     = RGBColor(0x0F, 0x17, 0x2A)   # near-black navy text
MUTED   = RGBColor(0x47, 0x55, 0x69)   # slate-600
FAINT   = RGBColor(0x94, 0xA3, 0xB8)   # slate-400
ACCENT  = RGBColor(0x25, 0x63, 0xEB)   # blue-600
GREEN   = RGBColor(0x05, 0x96, 0x69)
RED     = RGBColor(0xDC, 0x26, 0x26)
AMBER   = RGBColor(0xD9, 0x77, 0x06)
CARD    = RGBColor(0xF1, 0xF5, 0xF9)   # slate-100
CARDBDR = RGBColor(0xE2, 0xE8, 0xF0)   # slate-200
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)

EMU_IN = 914400
SW, SH = 13.333, 7.5

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


# ---- helpers ---------------------------------------------------------------
def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    return s


def _noline(shape):
    shape.line.fill.background()


def rect(s, l, t, w, h, fill=None, line=None, line_w=1.0, shadow=False, radius=False):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(l), Inches(t), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def text(s, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=4, line_spacing=1.0, wrap=True):
    """runs: list of paragraphs; each paragraph is list of (txt, size, color, bold, italic, font)."""
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after); p.space_before = Pt(0)
        p.line_spacing = line_spacing
        for (txt, size, color, bold, italic, font) in para:
            r = p.add_run(); r.text = txt
            r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
            r.font.color.rgb = color; r.font.name = font
    return tb


def R(txt, size=14, color=INK, bold=False, italic=False, font="Calibri"):
    return (txt, size, color, bold, italic, font)


def header(s, kicker, title, sub=None):
    rect(s, 0.55, 0.55, 0.13, 0.72, fill=ACCENT)  # accent bar
    y = 0.5
    if kicker:
        text(s, 0.85, y, 11.8, 0.32, [[R(kicker.upper(), 12, ACCENT, True)]])
        y += 0.34
    text(s, 0.83, y, 11.9, 0.8, [[R(title, 30, INK, True)]])
    if sub:
        text(s, 0.85, y + 0.62, 11.8, 0.5, [[R(sub, 14, MUTED)]])


def footer(s):
    idx = len(prs.slides._sldIdLst)  # auto: current slide number
    text(s, 0.85, 7.04, 9, 0.3, [[R("VAD in Production/Assembly Lines  ·  Güler & Türk  ·  CSE 496", 9, FAINT)]])
    text(s, 12.0, 7.04, 0.9, 0.3, [[R(str(idx), 9, FAINT)]], align=PP_ALIGN.RIGHT)


def add_video(s, video_path, poster_path, l, t, w, h, label, color):
    """Embed an mp4 (plays in PowerPoint) with a poster frame + a corner badge."""
    s.shapes.add_movie(str(video_path), Inches(l), Inches(t), Inches(w), Inches(h),
                       poster_frame_image=str(poster_path), mime_type="video/mp4")
    rect(s, l, t, w, h, fill=None, line=color, line_w=2.0)
    badge(s, l + 0.08, t + h - 0.36, 1.25, 0.28, label, color)


def pic_fit(s, path, l, t, w, h, border=None, border_w=1.5):
    """Place image fit inside box (contain), centered."""
    iw, ih = Image.open(path).size
    box_ar = w / h
    img_ar = iw / ih
    if img_ar > box_ar:
        nw = w; nh = w / img_ar
    else:
        nh = h; nw = h * img_ar
    nl = l + (w - nw) / 2
    nt = t + (h - nh) / 2
    p = s.shapes.add_picture(str(path), Inches(nl), Inches(nt), Inches(nw), Inches(nh))
    if border:
        b = rect(s, nl, nt, nw, nh, fill=None, line=border, line_w=border_w)
    return p


def pic_cover(s, path, l, t, w, h, border=None, border_w=1.5):
    """Crop image to fill box (cover)."""
    iw, ih = Image.open(path).size
    box_ar = w / h; img_ar = iw / ih
    pic = s.shapes.add_picture(str(path), Inches(l), Inches(t), Inches(w), Inches(h))
    if img_ar > box_ar:
        crop = (1 - box_ar / img_ar) / 2
        pic.crop_left = crop; pic.crop_right = crop
    else:
        crop = (1 - img_ar / box_ar) / 2
        pic.crop_top = crop; pic.crop_bottom = crop
    if border:
        rect(s, l, t, w, h, fill=None, line=border, line_w=border_w)
    return pic


def badge(s, l, t, w, h, label, color):
    b = rect(s, l, t, w, h, fill=color, radius=True)
    text(s, l, t, w, h, [[R(label, 10, WHITE, True)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def card(s, l, t, w, h, accent=ACCENT):
    rect(s, l, t, w, h, fill=CARD, line=CARDBDR, line_w=1, radius=True)
    rect(s, l, t, w, 0.09, fill=accent)  # top accent strip
    return l, t, w, h


def bullets(s, l, t, w, h, items, size=14, color=INK, gap=6, dot=ACCENT, lh=1.05):
    paras = []
    for it in items:
        if isinstance(it, tuple):
            txt, c, b = it
        else:
            txt, c, b = it, color, False
        paras.append([R("•  ", size, dot, True), R(txt, size, c, b)])
    text(s, l, t, w, h, paras, space_after=gap, line_spacing=lh)


# ===========================================================================
# SLIDE 1 — TITLE
# ===========================================================================
s = slide()
rect(s, 0, 0, SW, SH, fill=WHITE)
rect(s, 0, 0, 0.35, SH, fill=ACCENT)
rect(s, 0, 6.95, SW, 0.55, fill=CARD)
text(s, 1.0, 1.55, 11.5, 0.4, [[R("CSE 496 · GRADUATION PROJECT · GEBZE TECHNICAL UNIVERSITY", 13, ACCENT, True)]])
text(s, 0.97, 2.05, 11.6, 1.7,
     [[R("Video Anomaly Detection in", 40, INK, True)],
      [R("Production / Assembly Lines", 40, ACCENT, True)]], space_after=2)
text(s, 1.0, 3.95, 11.3, 0.8,
     [[R("A training-free, explainable approach using zero-shot Vision-Language Models", 18, MUTED)]])
# author / advisor cards
card(s, 1.0, 4.95, 5.3, 1.55)
text(s, 1.25, 5.15, 4.9, 0.4, [[R("AUTHORS", 11, ACCENT, True)]])
text(s, 1.25, 5.5, 4.9, 0.9,
     [[R("Alper Kaan Güler", 17, INK, True)], [R("Yunus Emre Türk", 17, INK, True)]], space_after=2)
card(s, 6.55, 4.95, 5.3, 1.55, accent=GREEN)
text(s, 6.8, 5.15, 4.9, 0.4, [[R("ADVISOR", 11, GREEN, True)]])
text(s, 6.8, 5.5, 4.9, 0.9,
     [[R("Asst. Prof. Habil Kalkan", 17, INK, True)],
      [R("Computer Engineering Dept.", 13, MUTED)]], space_after=2)

# ===========================================================================
# SLIDE — PROJECT DEFINITION
# ===========================================================================
s = slide()
header(s, "Overview", "Project Definition")
# one-line definition banner
rect(s, 0.85, 1.7, 11.6, 1.15, fill=RGBColor(0xEF, 0xF6, 0xFF), line=ACCENT, line_w=1.25, radius=True)
rect(s, 0.85, 1.7, 0.13, 1.15, fill=ACCENT)
text(s, 1.2, 1.9, 11.0, 0.8,
     [[R("We build a ", 15, INK), R("training-free, explainable system", 15, ACCENT, True),
       R(" that watches a production / assembly line and decides whether each video clip is ", 15, INK),
       R("normal or anomalous — and explains why.", 15, INK, True)]], line_spacing=1.1,
     anchor=MSO_ANCHOR.MIDDLE)
# four definition cards
defs = [
    ("Objective", "Detect line anomalies (wrong colour, orientation, missing part, belt stop, motion) without collecting or labelling any training data.", ACCENT),
    ("Approach", "Prompt a frozen Vision-Language Model (qwen3.6-plus) with the scene description, anomaly rules and learned guiding questions.", GREEN),
    ("Output", "A clip-level verdict — normal vs anomaly — plus the anomaly type and a short, human-readable reason for the decision.", AMBER),
    ("Key constraint", "Zero training, zero fine-tuning. The model adapts only through the prompt, so it re-targets to a new product via natural language.", RED),
]
gx, gy, cw, ch = 0.85, 3.1, 5.75, 1.7
for i, (tt, bb, acc) in enumerate(defs):
    r, c = divmod(i, 2)
    l = gx + c * (cw + 0.1); t = gy + r * (ch + 0.15)
    card(s, l, t, cw, ch, accent=acc)
    text(s, l + 0.28, t + 0.22, cw - 0.5, 0.4, [[R(tt, 15.5, INK, True)]])
    text(s, l + 0.28, t + 0.68, cw - 0.5, 0.95, [[R(bb, 12, MUTED)]], line_spacing=1.05)
footer(s)

# ===========================================================================
# SLIDE 2 — PROBLEM & MOTIVATION
# ===========================================================================
s = slide()
header(s, "Problem & Motivation", "Why classic computer vision struggles on the factory floor")
cards = [
    ("Data hunger", "Classic detectors need thousands of labelled defect images. In an optimised factory, real defects are rare — collecting them is slow and costly.", RED),
    ("Retraining cost", "Every product or layout change forces a full re-train. We want a system that adapts from a natural-language spec, not new training data.", AMBER),
    ("Black box", "Standard models flag an anomaly but cannot say WHY. Operators need a human-readable reason before they trust and act on an alarm.", ACCENT),
]
x = 0.85
for title, body, acc in cards:
    card(s, x, 1.75, 3.78, 2.55, accent=acc)
    text(s, x + 0.28, 2.0, 3.25, 0.5, [[R(title, 18, INK, True)]])
    text(s, x + 0.28, 2.55, 3.25, 1.7, [[R(body, 13, MUTED)]], line_spacing=1.05)
    x += 3.98
# our approach banner
rect(s, 0.85, 4.62, 11.6, 1.5, fill=RGBColor(0xEC, 0xFD, 0xF5), line=GREEN, line_w=1.25, radius=True)
rect(s, 0.85, 4.62, 0.13, 1.5, fill=GREEN)
text(s, 1.2, 4.82, 11.0, 0.45, [[R("OUR APPROACH — ZERO-SHOT PROMPTING", 13, GREEN, True)]])
text(s, 1.2, 5.25, 11.0, 0.85,
     [[R("Instead of training a detector, we prompt a ", 14, INK),
       R("frozen Vision-Language Model (qwen3.6-plus)", 14, INK, True),
       R(" with the scene description and the anomaly rules. Zero training, full explainability.", 14, INK)]],
     line_spacing=1.05)
footer(s)

# ===========================================================================
# SLIDE 3 — DATASET JOURNEY (IPAD -> why needed)
# ===========================================================================
s = slide()
header(s, "Dataset · why we built our own",
       "From a public benchmark to a dataset we can trust")
# Phase 1 IPAD
card(s, 0.85, 1.7, 4.7, 4.55, accent=FAINT)
text(s, 1.1, 1.92, 4.2, 0.35, [[R("PHASE 1 · STARTING POINT", 11, MUTED, True)]])
text(s, 1.1, 2.26, 4.2, 0.4, [[R("IPAD benchmark — button conveyor", 15, INK, True)]])
pic_cover(s, ASSETS / "ipad_button_1.jpg", 1.1, 2.75, 2.05, 1.5, border=CARDBDR)
pic_cover(s, ASSETS / "ipad_button_2.jpg", 3.28, 2.75, 2.05, 1.5, border=CARDBDR)
bullets(s, 1.1, 4.4, 4.3, 1.7, [
    "R01 push-button & S08 sorting scenes",
    "Anomalies: pin angle, cap colour, sorting / clog",
    ("Only ~256×256 — too coarse for subtle defects", RED, True),
], size=12, gap=5, dot=MUTED)
# arrow
ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(5.72), Inches(3.55), Inches(1.0), Inches(0.7))
ar.fill.solid(); ar.fill.fore_color.rgb = ACCENT; _noline(ar); ar.shadow.inherit = False
text(s, 5.5, 4.32, 1.45, 0.7, [[R("too low-res to judge colour & angle", 9, MUTED)]],
     align=PP_ALIGN.CENTER)
# Phase 2 ours
card(s, 6.9, 1.7, 5.55, 4.55, accent=ACCENT)
text(s, 7.15, 1.92, 5.0, 0.35, [[R("PHASE 2 · OUR OWN DATASET", 11, ACCENT, True)]])
text(s, 7.15, 2.26, 5.0, 0.4, [[R("Custom pen-conveyor rig", 15, INK, True)]])
fr = [("normal.jpg", "NORMAL", GREEN), ("anomaly_color.jpg", "COLOUR", RED), ("anomaly_orient.jpg", "ORIENT.", RED)]
fx = 7.15
for fn, lab, col in fr:
    pic_cover(s, ASSETS / fn, fx, 2.75, 1.62, 1.5, border=col, border_w=1.75)
    badge(s, fx + 0.05, 4.0, 0.95, 0.22, lab, col)
    fx += 1.72
bullets(s, 7.15, 4.42, 5.1, 1.7, [
    "13 clips · recorded in 4K → 720×1280 for the API",
    "7 controlled anomaly types · 3 identical pens · fixed camera",
    "Full control over what is normal vs anomalous",
], size=12, gap=5)
footer(s)

# ===========================================================================
# SLIDE 4 — PENS DATASET DETAIL (7 anomaly types)
# ===========================================================================
s = slide()
header(s, "Our dataset", "Seven controlled anomaly types on the pen line")
types = [
    ("normal.jpg",          "NORMAL",            GREEN, "Three identical pens, caps aligned, belt moving"),
    ("anomaly_color.jpg",   "WRONG COLOUR",      RED,   "One pen differs in colour from the others"),
    ("anomaly_orient.jpg",  "WRONG ORIENTATION", RED,   "A pen points the wrong way on the board"),
    ("anomaly_cap.jpg",     "MISSING CAP",       RED,   "A pen is missing its cap"),
    ("anomaly_pen.jpg",     "MISSING PEN",       RED,   "Fewer than three pens are present"),
    ("anomaly_temporal.jpg","TEMPORAL / MOTION", AMBER, "Pens shift mid-clip while belt looks normal"),
]
gx, gy = 0.85, 1.75
cw, ch = 3.78, 2.45
for i, (fn, lab, col, desc) in enumerate(types):
    r, c = divmod(i, 3)
    l = gx + c * (cw + 0.1)
    t = gy + r * (ch + 0.15)
    rect(s, l, t, cw, ch, fill=CARD, line=CARDBDR, line_w=1, radius=True)
    pic_cover(s, ASSETS / fn, l + 0.12, t + 0.12, cw - 0.24, 1.35, border=col, border_w=1.5)
    text(s, l + 0.15, t + 1.55, cw - 0.3, 0.32, [[R(lab, 12, col, True)]])
    text(s, l + 0.15, t + 1.86, cw - 0.3, 0.55, [[R(desc, 10.5, MUTED)]], line_spacing=1.0)
text(s, 0.85, 6.78, 11.6, 0.3,
     [[R("Belt-stop / freeze and a combined (orientation + missing-cap) case complete the 13 clips.", 11, FAINT, italic=True)]])
footer(s)

# ===========================================================================
# SLIDE — DATASET VIDEOS (embedded clips)
# ===========================================================================
s = slide()
header(s, "Our dataset · in motion", "Sample clips from the pen-conveyor rig")
text(s, 0.85, 1.62, 11.6, 0.45,
     [[R("Recorded in 4K, downscaled for the API. ", 13, INK),
       R("Click a clip to play it. ", 13, ACCENT, True),
       R("Same fixed camera, three pens — only the anomaly changes.", 13, INK)]])
vids = [
    ("normal.mp4",          "normal.jpg",         "NORMAL",   GREEN),
    ("wrong_color.mp4",     "anomaly_color.jpg",  "COLOUR",   RED),
    ("wrong_direction.mp4", "anomaly_orient.jpg", "ORIENT.",  RED),
    ("temporal_anomaly.mp4","anomaly_temporal.jpg","TEMPORAL", AMBER),
]
vw, vh = 2.55, 4.35
gap = (11.6 - 4 * vw) / 3.0
vx = 0.85
for fn, poster, lab, col in vids:
    add_video(s, ASSETS / "videos" / fn, ASSETS / poster, vx, 2.2, vw, vh, lab, col)
    vx += vw + gap
text(s, 0.85, 6.72, 11.6, 0.35,
     [[R("The temporal clip looks normal frame-by-frame — the anomaly is only in how the pens move over time.", 11, FAINT, italic=True)]])
footer(s)

# ===========================================================================
# SLIDE 5 — TWO APPROACHES (overview) + persona
# ===========================================================================
s = slide()
header(s, "Approach 1 (on IPAD buttons)", "First idea: robustness through multi-persona voting")
text(s, 0.85, 1.7, 11.6, 0.8,
     [[R("A single prompt was unstable. So many synthetic ", 14, INK),
       R("\"inspector\" personas", 14, ACCENT, True),
       R(" share the same rules but inspect the scene through a different lens, then ", 14, INK),
       R("vote", 14, INK, True), R(" — the majority decides.", 14, INK)]], line_spacing=1.05)
pcards = [
    ("Naïve operator", "Looks at obvious visuals — cap colour, which way the pins point."),
    ("Mechanical engineer", "Checks how the part sits and travels on the belt."),
    ("Vision engineer", "Pixel-level cues: subtle colour shift, angle, geometry."),
    ("Majority vote", "5 diverse voters + 1 consensus. A tie → LOW-confidence flag."),
]
x = 0.85
for i, (tt, bb) in enumerate(pcards):
    acc = ACCENT if i < 3 else GREEN
    card(s, x, 2.75, 2.85, 2.2, accent=acc)
    text(s, x + 0.22, 2.98, 2.45, 0.5, [[R(tt, 14.5, INK, True)]])
    text(s, x + 0.22, 3.5, 2.45, 1.3, [[R(bb, 12, MUTED)]], line_spacing=1.05)
    x += 2.95
rect(s, 0.85, 5.25, 11.6, 1.1, fill=RGBColor(0xFE, 0xF3, 0xC7), line=AMBER, line_w=1, radius=True)
rect(s, 0.85, 5.25, 0.13, 1.1, fill=AMBER)
text(s, 1.2, 5.42, 11.0, 0.85,
     [[R("Key finding:  ", 13, AMBER, True),
       R("compact YES/NO formats caused systematic false negatives. Forcing an ", 13, INK),
       R("\"articulate-first\"", 13, INK, True),
       R(" output (describe each frame, then decide) suppressed the model's bias toward \"normal\".", 13, INK)]],
     line_spacing=1.05)
footer(s)

# ===========================================================================
# SLIDE 6 — VERA PAPER BACKGROUND
# ===========================================================================
s = slide()
header(s, "The academic foundation", "VERA: Verbalized Learning for VAD (CVPR 2025)")
text(s, 0.85, 1.68, 11.6, 0.8,
     [[R("Ye, Liu & He — ", 14, INK), R("Explainable VAD via Verbalized Learning of VLMs.", 14, INK, italic=True),
       R("  Insight: a vague prompt gives only 53–65% AUC; concrete, ", 14, INK),
       R("learned guiding questions", 14, ACCENT, True), R(" do far better — with the model frozen.", 14, INK)]],
     line_spacing=1.05)
vc = [
    ("1 · Frozen VLM", "No fine-tuning, no extra reasoning module. The model is adapted only through the prompt.", ACCENT),
    ("2 · Verbalized learning", "Questions are learnable: a learner answers, an optimiser rewrites them from mistakes — using only coarse labels.", AMBER),
    ("3 · Coarse-to-fine", "Segment scores → scene-context ensemble → frame scores via temporal smoothing. SOTA on UCF-Crime & XD-Violence.", GREEN),
]
x = 0.85
for tt, bb, acc in vc:
    card(s, x, 2.75, 3.78, 2.35, accent=acc)
    text(s, x + 0.25, 2.98, 3.3, 0.5, [[R(tt, 16, INK, True)]])
    text(s, x + 0.25, 3.5, 3.3, 1.5, [[R(bb, 12.5, MUTED)]], line_spacing=1.05)
    x += 3.98
rect(s, 0.85, 5.45, 11.6, 0.95, fill=CARD, line=ACCENT, line_w=1, radius=True)
text(s, 1.15, 5.6, 11.1, 0.7,
     [[R("→  ", 14, ACCENT, True),
       R("We keep VERA's learner/optimiser question-learning core, but simplify it to ", 13.5, INK),
       R("clip-level verdicts", 13.5, INK, True), R(" for the factory setting  →  ", 13.5, INK),
       R("VERA-lite.", 13.5, ACCENT, True)]], anchor=MSO_ANCHOR.MIDDLE)
footer(s)

# ===========================================================================
# SLIDE 7 — VERA-LITE PIPELINE (step by step)
# ===========================================================================
s = slide()
header(s, "Approach 2 · our adaptation", "The VERA-lite pipeline, step by step")
steps = [
    ("1", "Guiding questions", "Start from concrete yes/no questions, one per anomaly family (count, colour, orientation, cap, belt motion, intra-object motion).", ACCENT),
    ("2", "Learner VLM", "The frozen model watches the video and answers every question YES/NO with the visual evidence, then emits a structured JSON verdict.", ACCENT),
    ("3", "Optimiser VLM", "It reads the learner's mistakes and rewrites the questions — sharper wording, no weights touched. We ran it once (v2 → v3).", AMBER),
    ("4", "Blind evaluation", "The learner NEVER sees the ground-truth label. An early leak was fixed; the first honest score was 0.77, then 0.92 after v3.", GREEN),
]
x = 0.85
cw = 2.85
for num, tt, bb, acc in steps:
    card(s, x, 1.85, cw, 3.2, accent=acc)
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.25), Inches(2.05), Inches(0.6), Inches(0.6))
    circ.fill.solid(); circ.fill.fore_color.rgb = acc; _noline(circ); circ.shadow.inherit = False
    text(s, x + 0.25, 2.05, 0.6, 0.6, [[R(num, 20, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.25, 2.8, cw - 0.5, 0.5, [[R(tt, 15, INK, True)]])
    text(s, x + 0.25, 3.3, cw - 0.5, 1.6, [[R(bb, 11.5, MUTED)]], line_spacing=1.05)
    x += cw + 0.1
rect(s, 0.85, 5.35, 11.6, 1.0, fill=RGBColor(0xEC, 0xFD, 0xF5), line=GREEN, line_w=1, radius=True)
text(s, 1.15, 5.5, 11.1, 0.75,
     [[R("Result of the loop:  ", 13.5, GREEN, True),
       R("rewriting the questions from negative constraints to positive verification lifted blind accuracy from ", 13.5, INK),
       R("0.77 → 0.92", 13.5, INK, True), R(".", 13.5, INK)]], anchor=MSO_ANCHOR.MIDDLE)
footer(s)

# ===========================================================================
# SLIDE 8 — GUIDING QUESTIONS v2 -> v3 (table)
# ===========================================================================
s = slide()
header(s, "Inside the optimiser", "How the guiding questions were hardened (v2 → v3)")
text(s, 0.85, 1.62, 11.6, 0.55,
     [[R("The optimiser noticed: ", 13, INK),
       R("\"VLMs struggle with logical negation.\"", 13, ACCENT, True),
       R("  It shifted prompts from negative constraints to positive verification.", 13, INK)]])
rows = [
    ("Anomaly family", "v2  (negated / weak)", "v3  (positive / hardened)"),
    ("Count", "no missing pen and no empty gap", "exactly three pens in a single row"),
    ("Colour", "none differing in colour", "identical barrel & cap colours, no mismatch"),
    ("Belt motion", "without stopping, pausing, reversing", "advances at every moment, never halting even briefly (1–3 s)"),
]
tl, tt, tw, th = 0.85, 2.35, 11.6, 3.0
tbl = s.shapes.add_table(len(rows), 3, Inches(tl), Inches(tt), Inches(tw), Inches(th)).table
tbl.columns[0].width = Inches(2.4); tbl.columns[1].width = Inches(4.3); tbl.columns[2].width = Inches(4.9)
for ci in range(3):
    for ri in range(len(rows)):
        cell = tbl.cell(ri, ci)
        cell.margin_left = Inches(0.12); cell.margin_right = Inches(0.1)
        cell.margin_top = Inches(0.06); cell.margin_bottom = Inches(0.06)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tfc = cell.text_frame; tfc.word_wrap = True
        p = tfc.paragraphs[0]; r = p.add_run(); r.text = rows[ri][ci]
        if ri == 0:
            cell.fill.solid(); cell.fill.fore_color.rgb = ACCENT
            r.font.color.rgb = WHITE; r.font.bold = True; r.font.size = Pt(13)
        else:
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE if ri % 2 else CARD
            r.font.size = Pt(12.5)
            if ci == 0:
                r.font.bold = True; r.font.color.rgb = INK
            elif ci == 1:
                r.font.color.rgb = MUTED
            else:
                r.font.color.rgb = ACCENT; r.font.bold = True
        r.font.name = "Calibri"
text(s, 0.85, 5.7, 11.6, 0.6,
     [[R("Positive, concrete phrasing removed the model's blind spots on count, colour and belt-stop cases.", 12.5, MUTED, italic=True)]])
footer(s)

# ===========================================================================
# SLIDE 9 — INSIDE A VERDICT (JSON)
# ===========================================================================
s = slide()
header(s, "Explainability", "Inside a verdict — structured, auditable JSON")
text(s, 0.85, 1.7, 4.7, 3.6,
     [[R("For every clip the model returns structured JSON: a YES/NO answer ", 14, INK),
       R("with the visual evidence", 14, INK, True),
       R(" for each guiding question, then the verdict, type, score and confidence.", 14, INK)]],
     line_spacing=1.1)
rect(s, 0.85, 3.5, 4.7, 1.9, fill=RGBColor(0xEC, 0xFD, 0xF5), line=GREEN, line_w=1, radius=True)
rect(s, 0.85, 3.5, 0.13, 1.9, fill=GREEN)
text(s, 1.15, 3.7, 4.25, 1.6,
     [[R("On the wrong-colour clip, a single ", 13, INK), R("\"NO\"", 13, RED, True),
       R(" on the colour question drives the decision. The per-question evidence makes the verdict ", 13, INK),
       R("auditable on the factory floor.", 13, INK, True)]], line_spacing=1.1)
# JSON panel
jl, jt, jw, jh = 5.9, 1.7, 6.55, 4.95
rect(s, jl, jt, jw, jh, fill=RGBColor(0x0F, 0x17, 0x2A), radius=True)
mono = "Consolas"
jlines = [
    [R("{", 12.5, FAINT, font=mono)],
    [R('  "question_answers": [', 12.5, RGBColor(0x93,0xC5,0xFD), font=mono)],
    [R('    {q1 ', 12.5, RGBColor(0xCB,0xD5,0xE1), font=mono), R('"YES"', 12.5, GREEN, True, font=mono), R('  three pens in a row}', 12.5, RGBColor(0xCB,0xD5,0xE1), font=mono)],
    [R('    {q2 ', 12.5, RGBColor(0xCB,0xD5,0xE1), font=mono), R('"YES"', 12.5, GREEN, True, font=mono), R('  caps aligned, same way}', 12.5, RGBColor(0xCB,0xD5,0xE1), font=mono)],
    [R('    {q3 ', 12.5, RGBColor(0xCB,0xD5,0xE1), font=mono), R('"NO"', 12.5, RED, True, font=mono), R('   leftmost pen is BLUE}', 12.5, RGBColor(0xFE,0xCA,0xCA), font=mono)],
    [R('    {q4 ', 12.5, RGBColor(0xCB,0xD5,0xE1), font=mono), R('"YES"', 12.5, GREEN, True, font=mono), R('  all capped, similar length}', 12.5, RGBColor(0xCB,0xD5,0xE1), font=mono)],
    [R('    {q5 ', 12.5, RGBColor(0xCB,0xD5,0xE1), font=mono), R('"YES"', 12.5, GREEN, True, font=mono), R('  board moves, no stop}', 12.5, RGBColor(0xCB,0xD5,0xE1), font=mono)],
    [R('    {q6 ', 12.5, RGBColor(0xCB,0xD5,0xE1), font=mono), R('"YES"', 12.5, GREEN, True, font=mono), R('  pens fixed to the board}', 12.5, RGBColor(0xCB,0xD5,0xE1), font=mono)],
    [R('  ],', 12.5, RGBColor(0x93,0xC5,0xFD), font=mono)],
    [R('  "verdict":      ', 12.5, RGBColor(0x93,0xC5,0xFD), font=mono), R('"ANOMALY"', 12.5, RED, True, font=mono)],
    [R('  "anomaly_type": ', 12.5, RGBColor(0x93,0xC5,0xFD), font=mono), R('"wrong_color"', 12.5, AMBER, font=mono)],
    [R('  "anomaly_score":', 12.5, RGBColor(0x93,0xC5,0xFD), font=mono), R(' 1.0', 12.5, AMBER, font=mono)],
    [R('  "confidence":   ', 12.5, RGBColor(0x93,0xC5,0xFD), font=mono), R('"HIGH"', 12.5, AMBER, font=mono)],
    [R("}", 12.5, FAINT, font=mono)],
]
text(s, jl + 0.3, jt + 0.25, jw - 0.5, jh - 0.4, jlines, space_after=2, line_spacing=1.0)
text(s, jl, jt + jh + 0.02, jw, 0.3,
     [[R("Verbatim learner output (blind v3 run) — abbreviated for display.", 9.5, FAINT, italic=True)]],
     align=PP_ALIGN.CENTER)
footer(s)

# ===========================================================================
# SLIDE 10 — THREE INPUT REPRESENTATIONS WE EXPLORED
# ===========================================================================
s = slide()
header(s, "Experiments", "Three ways we fed the clip to the model")

def _sq(l, t, sz, fill=RGBColor(0xDB, 0xEA, 0xFE), line=ACCENT, lw=1.25):
    rect(s, l, t, sz, sz, fill=fill, line=line, line_w=lw, radius=True)

reps = [
    ("Continuous video", "Send the whole clip; the VLM samples a few frames internally.",
     "12/13 · misses temporal", RED, ACCENT),
    ("Dense frames", "Extract N evenly-spaced frames, send them as separate images.",
     "11/13 · catches temporal", GREEN, ACCENT),
    ("Temporal grid", "Tile N time-ordered frames into ONE montage image, single pass.",
     "negative · ≈ dense + a false positive", RED, AMBER),
]
cw = 3.78
xs = [0.85, 4.73, 8.61]
for (tt, bb, chip, chipcol, acc), x in zip(reps, xs):
    card(s, x, 1.8, cw, 3.75, accent=acc)
    cx = x + cw / 2
    if tt.startswith("Continuous"):
        # timeline bar + play + ticks
        bw, bh = 2.7, 0.5
        bl, bt = cx - bw / 2, 2.55
        rect(s, bl, bt, bw, bh, fill=RGBColor(0xDB, 0xEA, 0xFE), line=ACCENT, line_w=1.25, radius=True)
        for k in range(1, 9):
            rect(s, bl + bw * k / 9.0, bt + 0.08, 0.012, bh - 0.16, fill=ACCENT)
        tri = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(cx - 0.22), Inches(2.95), Inches(0.44), Inches(0.5))
        tri.rotation = 90; tri.fill.solid(); tri.fill.fore_color.rgb = ACCENT; _noline(tri); tri.shadow.inherit = False
        text(s, x, 3.55, cw, 0.3, [[R("one continuous stream", 10.5, MUTED, italic=True)]], align=PP_ALIGN.CENTER)
    elif tt.startswith("Dense"):
        sz, gap = 0.62, 0.16
        tot = 4 * sz + 3 * gap
        sl = cx - tot / 2
        for k in range(4):
            _sq(sl + k * (sz + gap), 2.7, sz)
        text(s, x, 3.55, cw, 0.3, [[R("N discrete frames", 10.5, MUTED, italic=True)]], align=PP_ALIGN.CENTER)
    else:
        pic_fit(s, ASSETS / "grid_temporal.jpg", x + 0.3, 2.45, cw - 0.6, 1.45, border=AMBER, border_w=1.5)
        text(s, x, 3.92, cw, 0.3, [[R("3×3 montage we actually sent", 10, MUTED, italic=True)]], align=PP_ALIGN.CENTER)
    text(s, x + 0.25, 4.28, cw - 0.5, 0.4, [[R(tt, 16, INK, True)]])
    text(s, x + 0.25, 4.66, cw - 0.5, 0.7, [[R(bb, 11.5, MUTED)]], line_spacing=1.0)
    badge(s, x + 0.25, 5.18, cw - 0.5, 0.28, chip, chipcol)
# intrinsic trade-off banner
rect(s, 0.85, 5.78, 11.6, 0.95, fill=RGBColor(0xFE, 0xF3, 0xC7), line=AMBER, line_w=1, radius=True)
rect(s, 0.85, 5.78, 0.13, 0.95, fill=AMBER)
text(s, 1.15, 5.92, 11.1, 0.7,
     [[R("The trade-off is intrinsic:  ", 13, AMBER, True),
       R("continuous (video) catches edges & belt-stop; discrete (frames / grid) catches intra-object motion. We confirmed it on 3 discrete variants — ", 13, INK),
       R("no single representation reaches 13/13.", 13, INK, True)]], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0)
footer(s)

# ===========================================================================
# SLIDE 11 — REPRESENTATION TRADE-OFF & ENSEMBLE
# ===========================================================================
s = slide()
header(s, "The solution", "Disjoint blind spots → combine into an ensemble")
text(s, 0.85, 1.65, 7.0, 1.0,
     [[R("The hard case ", 13, INK, True),
       R("(temporal): pens shift mid-clip while the belt moves normally. ", 13, INK),
       R("Full-video input misses it — the model samples only a few internal frames. Dense frames catch it but miss some subtle cases. Their blind spots are ", 13, INK),
       R("disjoint", 13, ACCENT, True), R(".", 13, INK)]], line_spacing=1.05)
rows = [
    ("Input method", "Accuracy", "False pos.", "Temporal?"),
    ("Video (continuous)", "12/13 (0.92)", "0", "No"),
    ("Dense frames (16f)", "11/13 (0.85)", "0", "Yes"),
    ("OR-ensemble (video ∨ dense)", "13/13 (1.00)", "0", "Yes"),
]
tl, tt, tw, th = 0.85, 3.05, 7.05, 2.55
tbl = s.shapes.add_table(len(rows), 4, Inches(tl), Inches(tt), Inches(tw), Inches(th)).table
widths = [3.0, 1.55, 1.2, 1.3]
for i, wd in enumerate(widths):
    tbl.columns[i].width = Inches(wd)
for ri in range(len(rows)):
    for ci in range(4):
        cell = tbl.cell(ri, ci)
        cell.margin_left = Inches(0.1); cell.margin_top = Inches(0.04); cell.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]; r = p.add_run(); r.text = rows[ri][ci]
        r.font.name = "Calibri"
        if ri == 0:
            cell.fill.solid(); cell.fill.fore_color.rgb = INK
            r.font.color.rgb = WHITE; r.font.bold = True; r.font.size = Pt(12.5)
        else:
            last = (ri == len(rows) - 1)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0xDB,0xEA,0xFE) if last else (WHITE if ri % 2 else CARD)
            r.font.size = Pt(12.5); r.font.bold = last or ci == 0
            r.font.color.rgb = INK
            if ci == 3:
                r.font.color.rgb = GREEN if rows[ri][ci] == "Yes" else RED
                r.font.bold = True
            if last and ci == 1:
                r.font.color.rgb = ACCENT
text(s, 0.85, 5.75, 7.0, 0.7,
     [[R("Two views, disjoint blind spots → the union (OR) solves every clip with zero false alarms.", 12.5, MUTED, italic=True)]],
     line_spacing=1.0)
pic_fit(s, ASSETS / "fig_clip_matrix.png", 8.15, 1.95, 4.3, 4.6, border=CARDBDR)
text(s, 8.15, 6.45, 4.3, 0.3, [[R("Per-clip correctness: only the ensemble is fully green.", 10, FAINT, italic=True)]], align=PP_ALIGN.CENTER)
footer(s)

# ===========================================================================
# SLIDE 12 — RESULTS
# ===========================================================================
s = slide()
header(s, "Results", "Accuracy, speed and determinism")
# metric tiles
tiles = [("13/13", "Ensemble accuracy", GREEN), ("0", "False positives", ACCENT),
         ("0.846", "Repeatable single-view", INK), ("−41%", "Latency after tuning", AMBER)]
x = 0.85
for val, lab, col in tiles:
    card(s, x, 1.75, 2.85, 1.45, accent=col)
    text(s, x, 1.95, 2.85, 0.7, [[R(val, 30, col, True)]], align=PP_ALIGN.CENTER)
    text(s, x, 2.66, 2.85, 0.4, [[R(lab, 11.5, MUTED, True)]], align=PP_ALIGN.CENTER)
    x += 2.95
# left text
text(s, 0.85, 3.5, 5.6, 0.45, [[R("Speed & determinism", 17, INK, True)]])
bullets(s, 0.85, 3.95, 5.6, 2.6, [
    "qwen3.6-plus is a reasoning model — it once spent 15,684 reasoning tokens, spiking latency to 85.2 s/clip.",
    "Capping thinking_budget = 2000 → 49.9 s/clip (−41%).",
    ("Bonus: the cap made it deterministic. Verdict flips across 3 runs dropped from 2 clips to 0.", INK, True),
    "Honest, repeatable accuracy is 0.846 (11/13 every run) — not the lucky single-run 0.92.",
], size=12.5, gap=8)
# right figures
pic_fit(s, ASSETS / "fig_latency.png", 6.75, 3.45, 2.8, 2.95, border=CARDBDR)
pic_fit(s, ASSETS / "fig_accuracy.png", 9.65, 3.45, 2.8, 2.95, border=CARDBDR)
footer(s)

# ===========================================================================
# SLIDE 13 — LIMITATIONS (downsides, honest)
# ===========================================================================
s = slide()
header(s, "Honest assessment", "Limitations & open problems")
lims = [
    ("Small, single-scene dataset", "13 clips on one rig — results are indicative, not statistically conclusive. A larger, multi-scene set is needed."),
    ("LLM non-determinism", "The same input could flip the verdict; we needed a thinking-budget cap to stabilise it. True accuracy (0.846) is below the lucky 0.92."),
    ("No single view catches all", "Continuous video misses intra-object motion; dense frames miss subtle cases. The 13/13 needs an ensemble of two API calls."),
    ("Not real-time / cloud-bound", "~50 s per clip, depends on a paid API and internet. Raises latency, cost and data-privacy concerns for a real factory."),
    ("Limited model comparison", "Other models were inaccessible (qwen-vl-max-latest → 403, qwen3.7 → 400); results tie to one vendor model version."),
    ("Question tuning sees labels", "The optimiser used labels on a tiny set — risk of overfitting the questions to these specific clips."),
]
gx, gy, cw, ch = 0.85, 1.8, 5.75, 1.5
for i, (tt, bb) in enumerate(lims):
    r, c = divmod(i, 2)
    l = gx + c * (cw + 0.1); t = gy + r * (ch + 0.12)
    rect(s, l, t, cw, ch, fill=CARD, line=CARDBDR, line_w=1, radius=True)
    rect(s, l, t, 0.1, ch, fill=RED)
    text(s, l + 0.28, t + 0.13, cw - 0.5, 0.4, [[R(tt, 13.5, INK, True)]])
    text(s, l + 0.28, t + 0.55, cw - 0.5, 0.9, [[R(bb, 11, MUTED)]], line_spacing=1.0)
footer(s)

# ===========================================================================
# SLIDE 14 — CONCLUSION & FUTURE WORK
# ===========================================================================
s = slide()
rect(s, 0, 0, SW, SH, fill=WHITE)
rect(s, 0, 0, 0.35, SH, fill=ACCENT)
text(s, 1.0, 0.85, 11.5, 0.4, [[R("CONCLUSION", 13, ACCENT, True)]])
text(s, 0.97, 1.25, 11.6, 0.9,
     [[R("Explainable VAD without any labelled training data is viable for industry.", 26, INK, True)]],
     line_spacing=1.0)
# takeaways
card(s, 0.85, 2.55, 5.6, 3.45, accent=GREEN)
text(s, 1.1, 2.78, 5.1, 0.4, [[R("WHAT WE SHOWED", 12, GREEN, True)]])
bullets(s, 1.1, 3.25, 5.15, 2.6, [
    "A frozen VLM + learned guiding questions detects fine-grained line anomalies, zero training.",
    "Every verdict is a human-readable, auditable JSON.",
    "A video∨dense ensemble reaches 13/13 with 0 false alarms.",
    "A thinking-budget cap makes it 41% faster and deterministic.",
], size=12.5, gap=8, dot=GREEN)
# future
card(s, 6.55, 2.55, 5.9, 3.45, accent=ACCENT)
text(s, 6.8, 2.78, 5.4, 0.4, [[R("FUTURE WORK", 12, ACCENT, True)]])
bullets(s, 6.8, 3.25, 5.45, 2.6, [
    "Cascade: a cheap CV detector first, the VLM only on flagged segments → real-time & cheaper.",
    "Larger, multi-scene dataset for statistically solid numbers.",
    "On-prem / open VLM to remove cloud cost and privacy concerns.",
    "Automate the dense+video ensemble into one call.",
], size=12.5, gap=8)
text(s, 0.85, 6.25, 11.6, 0.5,
     [[R("Thank you — questions welcome.", 16, INK, True)]], align=PP_ALIGN.CENTER)
footer(s)

prs.save(str(OUT))
print("Saved:", OUT)
print("Slides:", len(prs.slides._sldIdLst))
